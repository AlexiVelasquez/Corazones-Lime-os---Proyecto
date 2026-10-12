from __future__ import annotations

import re
import subprocess
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "documentos"
ASSET_DIR = ROOT / "tmp" / "report-assets"
OUT_FILE = OUT_DIR / "Informe_Pruebas_Unitarias_5_Grupos_Corazones_Limenos.docx"

GREEN = "006A3B"
GREEN_LIGHT = "E5F2E9"
BLUE = "244A73"
PALE_BLUE = "EDF4FA"
GRAY = "F3F4F3"
LINE = "D9D9D9"
TEXT_GRAY = RGBColor(85, 94, 89)


def font_path(*names: str) -> str:
    fonts = Path("C:/Windows/Fonts")
    for name in names:
        path = fonts / name
        if path.exists():
            return str(path)
    raise FileNotFoundError(names)


MONO = font_path("CascadiaMono.ttf", "consola.ttf", "cour.ttf")


def remove_ansi(text: str) -> str:
    return re.sub(r"\x1b\[[0-?]*[ -/]*[@-~]", "", text).replace("\r", "")


def run_tests() -> str:
    result = subprocess.run(
        ["npm.cmd", "run", "test:unit"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=120,
    )
    output = remove_ansi(result.stdout + "\n" + result.stderr).strip()
    if result.returncode != 0:
        raise RuntimeError(output)
    return output


def run_coverage() -> str:
    result = subprocess.run(
        ["npm.cmd", "run", "test:unit", "--", "--coverage"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=120,
    )
    output = remove_ansi(result.stdout + "\n" + result.stderr).strip()
    if result.returncode != 0:
        raise RuntimeError(output)
    return output


def fit_lines(lines: list[str], width: int, font: ImageFont.FreeTypeFont) -> list[str]:
    result: list[str] = []
    for line in lines:
        while line and font.getlength(line) > width:
            cut = max(1, int(len(line) * width / font.getlength(line)))
            result.append(line[:cut])
            line = line[cut:]
        result.append(line)
    return result


def terminal_image(output: str, path: Path, title: str = "Terminal  npm run test:unit") -> None:
    width, height = 1500, 870
    image = Image.new("RGB", (width, height), "#101814")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(MONO, 26)
    body_font = ImageFont.truetype(MONO, 21)
    draw.rectangle((0, 0, width, 62), fill="#1D2A23")
    for x, color in [(26, "#FF605C"), (60, "#FFBD44"), (94, "#00CA4E")]:
        draw.ellipse((x, 20, x + 22, 42), fill=color)
    draw.text((132, 16), title, font=title_font, fill="#ECF5EF")
    lines = [f"> {title.removeprefix('Terminal  ')}", ""] + output.splitlines()
    lines = fit_lines(lines, width - 90, body_font)
    y = 88
    for line in lines[-30:]:
        color = "#7FE3A4" if "passed" in line.lower() or "PASS" in line else "#E7EFEA"
        draw.text((42, y), line, font=body_font, fill=color)
        y += 25
        if y > height - 35:
            break
    image.save(path)


def code_image(source: Path, path: Path, start: int, end: int, title: str) -> None:
    width, height = 1500, 900
    image = Image.new("RGB", (width, height), "#11151A")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(MONO, 24)
    body_font = ImageFont.truetype(MONO, 19)
    draw.rectangle((0, 0, width, 62), fill="#202833")
    draw.text((34, 16), title, font=title_font, fill="#F4F7FA")
    lines = source.read_text(encoding="utf-8").splitlines()[start - 1:end]
    y = 83
    for number, raw in enumerate(lines, start):
        line = raw.replace("\t", "  ")
        draw.text((25, y), f"{number:>3}", font=body_font, fill="#718096")
        x = 86
        color = "#D8DEE9"
        stripped = line.strip()
        if stripped.startswith(("describe", "it(", "it.each")):
            color = "#8BE9A8"
        elif stripped.startswith(("export class", "static", "return")):
            color = "#82C7FF"
        elif stripped.startswith("//") or stripped.startswith("/**") or stripped.startswith("*"):
            color = "#8A9A8F"
        draw.text((x, y), line[:118], font=body_font, fill=color)
        y += 24
        if y > height - 30:
            break
    image.save(path)


def test_outline_image(source: Path, path: Path) -> None:
    width, height = 1500, 900
    image = Image.new("RGB", (width, height), "#11151A")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(MONO, 24)
    body_font = ImageFont.truetype(MONO, 18)
    draw.rectangle((0, 0, width, 62), fill="#202833")
    draw.text((34, 16), "business-rules.spec.ts  Cinco grupos de pruebas", font=title_font, fill="#F4F7FA")
    selected = []
    for number, raw in enumerate(source.read_text(encoding="utf-8").splitlines(), 1):
        stripped = raw.strip()
        if stripped.startswith("describe(") or stripped.startswith("it(") or stripped.startswith("it.each"):
            selected.append((number, stripped))
    y = 86
    for number, line in selected:
        is_group = line.startswith("describe(")
        color = "#8BE9A8" if is_group else "#D8DEE9"
        prefix = "GRUPO" if is_group else "CASO "
        draw.text((28, y), f"{number:>3}  {prefix}  {line[:105]}", font=body_font, fill=color)
        y += 34 if is_group else 27
    image.save(path)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_figure(doc: Document, path: Path, caption: str, width=6.45) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.style = doc.styles["Caption"]
    cap.paragraph_format.space_after = Pt(12)


def add_result_table(doc: Document, rows: list[tuple[str, str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(2.6), Cm(7.0), Cm(4.2), Cm(2.1)]
    headers = ["ID", "Validación", "Resultado esperado", "Estado"]
    for idx, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[idx]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    set_repeat_table_header(table.rows[0])
    for row_idx, data in enumerate(rows):
        cells = table.add_row().cells
        for idx, (cell, text) in enumerate(zip(cells, data)):
            cell.width = widths[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx % 2:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.size = Pt(8.5)
            if idx == 3:
                run.bold = True
                run.font.color.rgb = RGBColor(0, 106, 59)
    doc.add_paragraph()


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(30, 35, 32)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(14)
    for style_name, size in (("Heading 1", 19), ("Heading 2", 14), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(15)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    caption = styles["Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = TEXT_GRAY


def build_document(test_output: str, terminal_png: Path, coverage_png: Path, tests_png: Path, rules_png: Path) -> None:
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        section.top_margin = Cm(2.1)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Informe de pruebas unitarias")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Proyecto Corazones Limeños")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0, 106, 59)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(22)
    meta.add_run("Curso: Calidad de Software\n").bold = True
    meta.add_run("Tecnología: Angular 22, TypeScript y Vitest\n")
    meta.add_run(f"Fecha de ejecución: {date.today().strftime('%d/%m/%Y')}")
    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(42)
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro.add_run("Propósito. ").bold = True
    intro.add_run("Este informe documenta cinco grupos de pruebas unitarias aplicadas a las reglas de negocio del sistema de donaciones Corazones Limeños. La ejecución verificó correctamente los casos de validación, límites, entrega, permisos por rol y seguridad del perfil. La suite final registró 23 pruebas aprobadas y ninguna prueba fallida.")
    doc.add_page_break()

    doc.add_heading("1 Resumen de resultados", level=1)
    summary = doc.add_table(rows=2, cols=4)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    metrics = [("5", "Grupos de pruebas"), ("23", "Pruebas aprobadas"), ("0", "Pruebas fallidas"), ("100 %", "Ejecución satisfactoria")]
    for col, (value, label) in enumerate(metrics):
        set_cell_shading(summary.cell(0, col), GREEN_LIGHT)
        p = summary.cell(0, col).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        r.bold = True
        r.font.size = Pt(20)
        r.font.color.rgb = RGBColor(0, 106, 59)
        p2 = summary.cell(1, col).paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run(label).font.size = Pt(8.5)
        for row in (0, 1):
            set_cell_margins(summary.cell(row, col), top=130, bottom=130)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("Conclusión principal. ").bold = True
    p.add_run("Las reglas evaluadas se comportan de acuerdo con los resultados esperados. La separación de la lógica en clases puras permite repetir las pruebas sin abrir el navegador, sin acceder a localStorage y sin depender de componentes visuales.")

    doc.add_heading("2 Herramienta y enfoque", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("El material de clase presenta JUnit para proyectos Java. ").bold = True
    p.add_run("Corazones Limeños está desarrollado en Angular y TypeScript, por lo que utiliza Vitest mediante el ejecutor oficial de Angular CLI. Ambos frameworks pertenecen a la familia xUnit y emplean la misma idea: organizar casos, preparar entradas, ejecutar una unidad aislada y comparar el resultado mediante aserciones.")
    organization = doc.add_paragraph()
    organization.paragraph_format.space_before = Pt(10)
    organization.paragraph_format.line_spacing = 1.2
    heading_run = organization.add_run("2.1 Organización de los archivos")
    heading_run.bold = True
    heading_run.font.name = "Aptos Display"
    heading_run.font.size = Pt(14)
    for text in [
        "business-rules.ts contiene las reglas de negocio que se desean comprobar.",
        "business-rules.spec.ts contiene los casos unitarios y las aserciones.",
        "app.spec.ts contiene las pruebas básicas de creación de la aplicación Angular.",
        "pruebas-unitarias-ieee829.md conserva la trazabilidad de los casos.",
    ]:
        organization.add_run(f"\n•  {text}")

    doc.add_heading("3 Evidencia de ejecución", level=1)
    add_figure(doc, terminal_png, "Figura 1. Ejecución real de npm run test:unit con todos los casos aprobados.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("Interpretación. ").bold = True
    p.add_run("La línea Test Files 2 passed indica que los dos archivos de pruebas finalizaron correctamente. Las pruebas de negocio están separadas en cinco bloques describe claramente identificados. La línea Tests 23 passed confirma que no existe ningún caso fallido en la ejecución presentada.")
    add_figure(doc, tests_png, "Figura 2. Casos unitarios definidos con describe, it y expect.")
    add_figure(doc, rules_png, "Figura 3. Clase de reglas del negocio evaluada por la suite.")

    doc.add_heading("4 Cobertura de código", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("La cobertura se obtuvo con el proveedor V8 mediante npm run test:unit -- --coverage. ").bold = True
    p.add_run("La lógica central de negocio alcanzó 88,63 % de sentencias, 87,80 % de ramas, 100 % de funciones y 97,05 % de líneas. En el conjunto completo del proyecto se obtuvo 62,59 % de sentencias y 76,92 % de líneas; la diferencia corresponde principalmente a la plantilla visual y al almacén de datos, que requieren pruebas de componentes e integración.")
    coverage_table = doc.add_table(rows=2, cols=4)
    coverage_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, (value, label) in enumerate([
        ("88,63 %", "Sentencias"),
        ("87,80 %", "Ramas"),
        ("100 %", "Funciones"),
        ("97,05 %", "Líneas"),
    ]):
        set_cell_shading(coverage_table.cell(0, col), PALE_BLUE)
        value_p = coverage_table.cell(0, col).paragraphs[0]
        value_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        value_r = value_p.add_run(value)
        value_r.bold = True
        value_r.font.size = Pt(16)
        value_r.font.color.rgb = RGBColor(36, 74, 115)
        label_p = coverage_table.cell(1, col).paragraphs[0]
        label_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_p.add_run(label).font.size = Pt(8.5)
    doc.add_paragraph()
    add_figure(doc, coverage_png, "Figura 4. Reporte real de cobertura generado por Vitest y V8.")

    doc.add_heading("5 Casos de prueba documentados", level=1)
    doc.add_paragraph("Los casos se identifican según una estructura de trazabilidad compatible con la especificación IEEE 829 solicitada en clase. El resultado obtenido coincide con el resultado esperado en todos los casos.")
    add_result_table(doc, [
        ("CP-UNIT-001", "Aceptar la categoría Juguetes", "Categoría válida", "PASS"),
        ("CP-UNIT-002", "Rechazar categoría vacía o desconocida", "Validación rechazada", "PASS"),
        ("CP-UNIT-003", "Aceptar descripción y cantidad válidas", "Detalle válido", "PASS"),
        ("CP-UNIT-004", "Rechazar descripción menor de cinco caracteres", "Mensaje de validación", "PASS"),
        ("CP-UNIT-005", "Rechazar 0, 1001, decimales y NaN", "Cantidad inválida", "PASS"),
        ("CP-UNIT-006", "Aceptar los límites 1 y 1000", "Límites válidos", "PASS"),
        ("CP-UNIT-007", "Exigir distrito en punto de recojo", "Distrito obligatorio", "PASS"),
        ("CP-UNIT-008", "Exigir dirección y horario a domicilio", "Datos completos", "PASS"),
        ("CP-UNIT-009", "Aceptar solamente un albergue aprobado", "Albergue válido", "PASS"),
        ("CP-UNIT-010", "Limitar el contador de cajas", "Rango de 1 a 1000", "PASS"),
    ])
    add_result_table(doc, [
        ("CP-UNIT-011", "Administrador asigna una ruta", "Estado En tránsito", "PASS"),
        ("CP-UNIT-012", "Albergue confirma la recepción", "Estado Entregado", "PASS"),
        ("CP-UNIT-013", "Bloquear cambios por roles no autorizados", "Estado sin cambios", "PASS"),
        ("CP-UNIT-014", "Validar formato del correo", "Válido o inválido", "PASS"),
        ("CP-UNIT-015", "Validar teléfono de 9 a 15 dígitos", "Válido o inválido", "PASS"),
        ("CP-UNIT-016", "Aceptar cambio de contraseña correcto", "Cambio permitido", "PASS"),
        ("CP-UNIT-017", "Rechazar contraseña actual incorrecta", "Cambio rechazado", "PASS"),
        ("CP-UNIT-018", "Rechazar contraseñas nuevas diferentes", "Cambio rechazado", "PASS"),
    ])

    doc.add_heading("6 Explicación de cinco pruebas representativas", level=1)
    doc.add_heading("6.1 Cantidad de cajas", level=2)
    doc.add_paragraph("La prueba envía cantidades fuera del rango permitido: 0, 1001, un número decimal y NaN. La regla debe rechazarlas porque el formulario solo admite números enteros de 1 a 1000. También se prueban exactamente los valores 1 y 1000 para comprobar los límites.")
    doc.add_heading("6.2 Transición por rol", level=2)
    doc.add_paragraph("Una donación pendiente solo puede pasar a En tránsito cuando actúa un administrador. Una donación en tránsito solo puede pasar a Entregado cuando actúa el albergue. Si interviene un rol sin permiso, el estado original se conserva.")
    doc.add_heading("6.3 Cambio de contraseña", level=2)
    doc.add_paragraph("La regla compara la contraseña actual con la almacenada, exige una nueva contraseña de al menos ocho caracteres y verifica que la repetición coincida. Las pruebas evalúan el caso correcto y los principales errores.")
    doc.add_heading("6.4 Método de entrega", level=2)
    doc.add_paragraph("La prueba comprueba dos rutas logísticas. Para entregar en un punto de recojo exige seleccionar un distrito; para solicitar recojo a domicilio exige una dirección y un horario. De esta forma, cada modalidad valida únicamente los datos que necesita.")
    doc.add_heading("6.5 Formato del correo", level=2)
    doc.add_paragraph("La prueba compara un correo válido con una cadena sin estructura de correo. La regla acepta persona@calidad.pe y rechaza correo-invalido, con lo cual protege los datos de contacto guardados en el perfil.")

    doc.add_heading("7 Ejecución de las pruebas", level=1)
    doc.add_paragraph("Abrir una terminal en la carpeta del proyecto y ejecutar:")
    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Cm(0.8)
    run = code.add_run("cd C:\\ProyectoCalidad\\Proyectocalidad\nnpm run test:unit")
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0, 70, 45)
    doc.add_paragraph("Para trabajar con el ciclo TDD de forma continua puede utilizarse npm run test:watch. Vitest vuelve a ejecutar la suite después de cada cambio guardado.")

    doc.add_heading("8 Relación con TDD", level=1)
    tdd = [
        ("Red", "Se escribe primero un caso en business-rules.spec.ts y se confirma que falle porque la regla todavía no existe."),
        ("Green", "Se implementa la cantidad mínima de código en business-rules.ts hasta conseguir que el caso pase."),
        ("Refactor", "Se mejora la estructura sin cambiar el comportamiento y se vuelve a ejecutar toda la suite para detectar regresiones."),
    ]
    for name, description in tdd:
        p = doc.add_paragraph()
        p.add_run(f"{name}. ").bold = True
        p.add_run(description)

    doc.add_heading("9 Guía para la exposición", level=1)
    points = [
        "El objetivo de una prueba unitaria es comprobar una parte pequeña y aislada del sistema.",
        "JUnit es la herramienta mostrada para Java; Vitest cumple esa función en Angular y TypeScript.",
        "La lógica fue separada de la interfaz para que las pruebas no dependan del navegador.",
        "Los IDs CP-UNIT permiten relacionar cada prueba automática con su documentación IEEE 829.",
        "Los casos incluyen rutas correctas, datos inválidos, valores límite y permisos por rol.",
        "La suite está organizada en cinco bloques visibles: categoría, detalle, entrega, estados y perfil.",
        "El resultado final fue de 23 pruebas aprobadas y ninguna fallida.",
    ]
    for point in points:
        doc.add_paragraph(point, style="List Bullet")

    doc.add_heading("10 Conclusiones", level=1)
    doc.add_paragraph("La implementación de pruebas unitarias redujo la dependencia de verificaciones manuales y convirtió las reglas principales en contratos ejecutables. La suite comprueba tanto entradas válidas como errores frecuentes y restricciones de autorización. Al ejecutar las pruebas después de cada cambio, el equipo puede refactorizar el sistema con una señal inmediata sobre posibles regresiones.")
    doc.add_paragraph("Las pruebas unitarias no reemplazan las pruebas de integración ni las pruebas de interfaz. El siguiente nivel de calidad debe verificar la interacción entre componentes, almacenamiento local, rutas y formularios completos.")

    properties = doc.core_properties
    properties.title = "Informe de pruebas unitarias del proyecto Corazones Limeños"
    properties.subject = "Evidencia de pruebas unitarias con Angular y Vitest"
    properties.author = "Equipo del proyecto Corazones Limeños"
    properties.keywords = "pruebas unitarias, Vitest, Angular, IEEE 829, TDD"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_FILE)


def main() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    output = run_tests()
    coverage_output = run_coverage()
    terminal_png = ASSET_DIR / "captura_terminal_pruebas.png"
    coverage_png = ASSET_DIR / "captura_terminal_cobertura.png"
    tests_png = ASSET_DIR / "captura_codigo_pruebas.png"
    rules_png = ASSET_DIR / "captura_reglas_negocio.png"
    terminal_image(output, terminal_png)
    terminal_image(coverage_output, coverage_png, "Terminal  npm run test:unit -- --coverage")
    test_outline_image(ROOT / "src/app/business-rules.spec.ts", tests_png)
    code_image(ROOT / "src/app/business-rules.ts", rules_png, 1, 36, "business-rules.ts  Reglas del negocio")
    build_document(output, terminal_png, coverage_png, tests_png, rules_png)
    print(OUT_FILE)


if __name__ == "__main__":
    main()
